"""Build an Excel snapshot of the filtered follow-up queue."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def workbook_bytes(rows):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Follow-ups"
    details = workbook.create_sheet("Attempts")
    summary.append(["Client", "Phone", "Rep", "Company", "Status", "Priority", "Created (UTC)",
                    "Calls completed", "Calls overdue", "Messages completed", "Messages overdue", "Lead ID"])
    details.append(["Lead ID", "Client", "Type", "Attempt", "Status", "Scheduled start (UTC)",
                    "Scheduled end (UTC)", "Completed (UTC)"])
    for row in rows:
        summary.append([row.get(key, "") for key in (
            "client", "client_phone", "rep", "company", "status", "priority", "created_at",
            "completed_count", "overdue_count", "completed_message_count", "overdue_message_count", "lead_id")])
        for attempt in row.get("timeline", []):
            details.append([row.get("lead_id", ""), row.get("client", ""),
                            *[attempt.get(key, "") for key in ("kind", "label", "status", "scheduled_start", "scheduled_end", "completed_at")]])
    for sheet in (summary, details):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
        for row in sheet:
            for cell in row:
                # Export source text literally, never as Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="032D60")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
